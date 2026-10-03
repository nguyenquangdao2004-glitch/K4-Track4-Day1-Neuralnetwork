"""results_table.py — Lưu trữ kết quả thí nghiệm ra JSON và xuất bảng Excel experiments.xlsx.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có.
    """
    p_dir = Path(results_dir)
    p_dir.mkdir(parents=True, exist_ok=True)

    exp_id = result["cfg"].get("exp_id", "unnamed_exp")
    file_path = p_dir / f"{exp_id}.json"

    data_to_save = {
        "cfg": result.get("cfg", {}),
        "history": result.get("history", {}),
        "summary": result.get("summary", {}),
    }

    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    print(f"Đã lưu kết quả thí nghiệm: {file_path}")
    return str(file_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p_dir = Path(results_dir)
    if not p_dir.exists():
        return []

    results = []
    for f in sorted(p_dir.glob("*.json")):
        try:
            with open(f, "r", encoding="utf-8") as fp:
                data = json.load(fp)
                results.append(data)
        except Exception as e:
            print(f"Lỗi khi đọc file {f}: {e}")

    results.sort(key=lambda r: r.get("cfg", {}).get("exp_id", ""))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng.
    """
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id = cfg.get("exp_id", "")

    eval_acc = eval_scores.get("accuracy") if eval_scores else ""
    eval_macro_f1 = eval_scores.get("macro_f1") if eval_scores else ""

    hidden_val = cfg.get("hidden", (256, 128))
    if isinstance(hidden_val, (list, tuple)):
        hidden_str = str(tuple(hidden_val))
    else:
        hidden_str = str(hidden_val)

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", "sgd_momentum"),
        "lr": cfg.get("lr", ""),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", "") if cfg.get("clip_norm") is not None else "",
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", ""),
        "best_val_loss": summary.get("best_val_loss", ""),
        "best_epoch": summary.get("best_epoch", ""),
        "final_train_loss": summary.get("final_train_loss", ""),
        "final_val_loss": summary.get("final_val_loss", ""),
        "val_acc": summary.get("val_acc", ""),
        "val_macro_f1": summary.get("val_macro_f1", ""),
        "time_per_epoch_s": summary.get("time_per_epoch_s", ""),
        "peak_mem_MB": summary.get("peak_mem_MB", ""),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_acc if eval_acc != "" else "",
        "eval_macro_f1": eval_macro_f1 if eval_macro_f1 != "" else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    import openpyxl

    tmpl_p = Path(template_path)
    out_p = Path(out_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)

    if not tmpl_p.exists():
        raise FileNotFoundError(f"Không tìm thấy template excel tại: {tmpl_p}")

    wb = openpyxl.load_workbook(str(tmpl_p))
    if "Experiments" not in wb.sheetnames:
        raise ValueError("Sheet 'Experiments' không tồn tại trong template excel!")

    ws = wb["Experiments"]

    # Đọc headers từ hàng 1
    col_mapping = {}
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col_idx).value
        if val:
            col_mapping[str(val).strip()] = col_idx

    formula_columns = {
        "step0_gap_vs_lnC",
        "gap_val_minus_train",
        "delta_val_f1_vs_base",
        "beyond_noise",
    }

    start_row = 2
    for r_idx, row_data in enumerate(rows, start=start_row):
        for key, value in row_data.items():
            if key in formula_columns:
                continue  # Bỏ qua các cột công thức để Excel tự tính
            if key in col_mapping:
                target_col = col_mapping[key]
                ws.cell(row=r_idx, column=target_col, value=value)

    wb.save(str(out_p))
    print(f"Đã lưu bảng kết quả: {out_p} (gồm {len(rows)} thí nghiệm)")
